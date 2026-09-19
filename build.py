#!/usr/bin/env python3
"""
CF Survivor League site builder.

Reads a season tracking workbook and emits data/season<NN>.json in the
schema the site renders from. Every season uses the same schema, so the
archive and the live season share one set of rendering code.

Usage:
    python3 build.py 51
    python3 build.py 51 --workbook "/path/to/Survivor Tracking Season 51.xlsx"
    python3 build.py all          # rebuild every season listed in SEASONS

The workbook is the source of truth. Nothing here writes back to it.
"""

import argparse
import json
import os
import sys
from datetime import date

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip3 install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(HERE, "data")

# Default location of the season folders. Override with --root.
DEFAULT_ROOT = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-CornerstoneFellowshipofLivermoreCalifornia/"
    "CF Work Files/11_CLAUDE-Workspace/Survivor"
)
# On the Cowork device bridge the same folder is mounted here.
FALLBACK_ROOTS = [os.path.expanduser("~/mnt/Survivor")]

SEASONS = [48, 49, 50, 51]

# Scoring categories that mean a castaway left the game.
ELIMINATION_LABELS = {
    "Voted Out WITH Idol",
    "Lose Vote",
    "Med Visit EVAC",
    "Quit Game",
}
WINNER_LABEL = "Win Survivor"

# Header rows inside each episode tab.
LABEL_ROW = 2      # scoring category names
VALUE_ROW = 3      # point value for each category
FIRST_DATA_ROW = 4

# Dashboard layout.
DASH_HEADER_ROW = 3
DASH_FIRST_ROW = 4
DASH_NAME_COL = 2   # B
DASH_SCORE_COL = 3  # C
DASH_PICK_COLS = (4, 5, 6)  # D, E, F


def find_workbook(season, root=None):
    roots = [root] if root else [DEFAULT_ROOT] + FALLBACK_ROOTS
    name = f"Survivor Tracking Season {season}.xlsx"
    for base in roots:
        path = os.path.join(base, f"Season {season}", name)
        if os.path.exists(path):
            return path
    raise FileNotFoundError(
        f"Could not find '{name}'. Looked under: {', '.join(roots)}"
    )


def cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def is_checked(value):
    """Excel checkboxes come through as booleans or as the strings TRUE/FALSE."""
    if isinstance(value, bool):
        return value
    return cell_text(value).upper() == "TRUE"


def episode_sheets(wb):
    """Episode tabs in order: E1, E2, ... however many the workbook has."""
    found = []
    for name in wb.sheetnames:
        stripped = name.strip()
        if len(stripped) > 1 and stripped[0].upper() == "E" and stripped[1:].isdigit():
            found.append((int(stripped[1:]), name))
    found.sort()
    return found


def read_scoring(ws):
    """Map each column to its scoring label and point value."""
    columns = []
    for col in range(1, ws.max_column + 1):
        label = cell_text(ws.cell(row=LABEL_ROW, column=col).value)
        raw = ws.cell(row=VALUE_ROW, column=col).value
        if not label or label == "NA":
            continue
        if not isinstance(raw, (int, float)):
            continue
        columns.append({"col": col, "label": label, "points": int(raw)})
    return columns


def read_castaways(wb, values):
    """
    Castaway roster. Prefers the Cast tab when the workbook has one
    (Season 51 onward); otherwise reads the names off the E1 tab.
    """
    roster = []
    if "Cast" in wb.sheetnames:
        ws = values["Cast"]
        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            name = cell_text(ws.cell(row=row, column=1).value)
            if not name:
                continue
            roster.append({
                "name": name,
                "tribe": cell_text(ws.cell(row=row, column=2).value),
            })
        return roster

    episodes = episode_sheets(wb)
    if not episodes:
        return roster
    ws = values[episodes[0][1]]
    tribe = ""
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        name = cell_text(ws.cell(row=row, column=3).value)
        if not name:
            continue
        found = cell_text(ws.cell(row=row, column=1).value)
        if found:
            tribe = found
        roster.append({"name": name, "tribe": tribe})
    return roster


def build_season(season, workbook_path, config):
    wb = openpyxl.load_workbook(workbook_path, data_only=False)
    wb_values = openpyxl.load_workbook(workbook_path, data_only=True)
    values = {name: wb_values[name] for name in wb_values.sheetnames}

    season_cfg = config.get("seasons", {}).get(str(season), {})
    episodes_meta = {int(k): v for k, v in season_cfg.get("episodes", {}).items()}

    roster = read_castaways(wb, values)
    by_name = {c["name"]: c for c in roster}
    for castaway in roster:
        castaway.update({
            "status": "IN",
            "points": 0,
            "out_episode": None,
            "out_reason": None,
            "winner": False,
            "events": [],
            "drafted_by": [],
        })

    scoring_table = None
    episodes = []

    for number, sheet_name in episode_sheets(wb):
        ws = values[sheet_name]
        columns = read_scoring(ws)
        if scoring_table is None and columns:
            scoring_table = columns

        scored = False
        episode_events = []

        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            name = cell_text(ws.cell(row=row, column=3).value)
            if not name or name not in by_name:
                continue
            castaway = by_name[name]
            for column in columns:
                if not is_checked(ws.cell(row=row, column=column["col"]).value):
                    continue
                scored = True
                castaway["points"] += column["points"]
                castaway["events"].append({
                    "episode": number,
                    "label": column["label"],
                    "points": column["points"],
                })
                episode_events.append({
                    "castaway": name,
                    "label": column["label"],
                    "points": column["points"],
                })
                if column["label"] in ELIMINATION_LABELS and castaway["status"] == "IN":
                    castaway["status"] = "OUT"
                    castaway["out_episode"] = number
                    castaway["out_reason"] = column["label"]
                if column["label"] == WINNER_LABEL:
                    castaway["winner"] = True

        meta = episodes_meta.get(number, {})
        episodes.append({
            "number": number,
            "label": f"E{number}",
            "title": meta.get("title", ""),
            "airs": meta.get("airs"),
            "scored": scored,
            "events": episode_events,
        })

    players = []
    if "Dashboard" in values:
        ws = values["Dashboard"]
        for row in range(DASH_FIRST_ROW, ws.max_row + 1):
            name = cell_text(ws.cell(row=row, column=DASH_NAME_COL).value)
            if not name:
                continue
            picks = [
                cell_text(ws.cell(row=row, column=col).value)
                for col in DASH_PICK_COLS
            ]
            picks = [p for p in picks if p]
            # Scores are computed from the episode tabs rather than read from
            # the Dashboard's cached formula results. Excel only refreshes
            # those cached values when the file is opened, so anything that
            # edits the workbook without launching Excel would otherwise
            # produce a site showing stale numbers.
            per_episode = {}
            for episode in episodes:
                per_episode[episode["label"]] = sum(
                    event["points"]
                    for event in episode["events"]
                    if event["castaway"] in picks
                )
            total = sum(per_episode.values())
            players.append({
                "name": name,
                "picks": picks,
                "total": total,
                "episodes": per_episode,
            })
            for pick in picks:
                if pick in by_name:
                    by_name[pick]["drafted_by"].append(name)

    # Rank by total, then alphabetically. Ties share a rank.
    players.sort(key=lambda p: (-p["total"], p["name"].lower()))
    last_total = None
    last_rank = 0
    for index, player in enumerate(players, start=1):
        if player["total"] != last_total:
            last_rank = index
            last_total = player["total"]
        player["rank"] = last_rank

    # Movement versus the previous scored episode.
    scored_labels = [e["label"] for e in episodes if e["scored"]]
    if len(scored_labels) >= 1:
        previous = {}
        for player in players:
            through = sum(
                player["episodes"].get(label, 0) for label in scored_labels[:-1]
            )
            previous[player["name"]] = through
        ordered = sorted(
            players, key=lambda p: (-previous[p["name"]], p["name"].lower())
        )
        prior_rank = {p["name"]: i for i, p in enumerate(ordered, start=1)}
        for player in players:
            player["previous_rank"] = prior_rank[player["name"]]
            player["movement"] = prior_rank[player["name"]] - player["rank"]
    else:
        for player in players:
            player["previous_rank"] = player["rank"]
            player["movement"] = 0

    positives = [c for c in (scoring_table or []) if c["points"] > 0]
    negatives = [c for c in (scoring_table or []) if c["points"] < 0]

    if not players:
        status = "pre_draft"
    elif any(c["winner"] for c in roster):
        status = "complete"
    else:
        status = "in_progress"

    return {
        "season": season,
        "name": season_cfg.get("name", f"Survivor {season}"),
        "subtitle": season_cfg.get("subtitle", ""),
        "status": status,
        "generated": date.today().isoformat(),
        "source_workbook": os.path.basename(workbook_path),
        "tribes": season_cfg.get("tribes", []),
        "castaways": roster,
        "players": players,
        "episodes": episodes,
        "scoring": {
            "positive": [
                {"label": c["label"], "points": c["points"]} for c in positives
            ],
            "negative": [
                {"label": c["label"], "points": c["points"]} for c in negatives
            ],
        },
    }


def load_config():
    path = os.path.join(DATA_DIR, "config.json")
    if not os.path.exists(path):
        return {}
    with open(path) as handle:
        return json.load(handle)


def main():
    parser = argparse.ArgumentParser(description="Build CF Survivor season data.")
    parser.add_argument("season", help="season number, or 'all'")
    parser.add_argument("--workbook", help="explicit path to the tracking workbook")
    parser.add_argument("--root", help="folder holding the Season NN subfolders")
    args = parser.parse_args()

    config = load_config()
    os.makedirs(DATA_DIR, exist_ok=True)

    targets = SEASONS if args.season == "all" else [int(args.season)]
    for season in targets:
        try:
            path = args.workbook or find_workbook(season, args.root)
        except FileNotFoundError as error:
            print(f"  skip S{season}: {error}")
            continue
        data = build_season(season, path, config)
        out = os.path.join(DATA_DIR, f"season{season}.json")
        with open(out, "w") as handle:
            json.dump(data, handle, indent=1)
        scored = sum(1 for e in data["episodes"] if e["scored"])
        print(
            f"  S{season}: {len(data['players'])} players, "
            f"{len(data['castaways'])} castaways, {scored} episodes scored "
            f"-> {os.path.relpath(out, HERE)}"
        )


if __name__ == "__main__":
    main()
